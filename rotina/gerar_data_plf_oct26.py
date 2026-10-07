#!/usr/bin/env python3
"""Gera docs/data-plf-oct26.json (contrato fortraffic-dashboard/1) para Familias Estelares PLF OCT26.

Entradas (as mesmas que a rotina já baixa):
  /tmp/leads.xlsx      aba OCT26        [0]FECHA [1]NOMBRE [2]EMAIL [3]utm_source [4]utm_medium [5]utm_term [6]utm_content
  /tmp/pesquisa.xlsx   aba "Respuestas de formulario 1"
  /tmp/meta2/ads.json        level ad, time_increment 1: name, campaign_name, adset_name, effective_status, amount_spent, impressions, link_click, landing_page_view
  /tmp/meta2/idadegen.json   level adset, time_increment 1, breakdowns [age, gender]: campaign_name, amount_spent, lead
  /tmp/meta2/posicao.json    level adset, time_increment 1, breakdowns [publisher_platform, platform_position]: campaign_name, amount_spent, lead
  /tmp/leitura.json          {"pt": "...", "es": "..."} (opcional; HTML simples)
Saída: /tmp/data-plf-oct26.json
"""
import json, re, sys, collections
from datetime import datetime, timedelta
from zoneinfo import ZoneInfo
import pandas as pd
sys.path.insert(0, __import__('os').path.dirname(__import__('os').path.abspath(__file__)))
import build_data_json as B

BASE = '/tmp'
AGORA = datetime.now(ZoneInfo('America/Sao_Paulo'))
PAGOS = ('ig', 'fb', 'an', 'th')
TESTE = re.compile(r'prueba|^test|\btest\b', re.I)


def meta_rows(nome):
    try:
        o = json.load(open(f'{BASE}/meta2/{nome}.json', encoding='utf-8'))
    except FileNotFoundError:
        return []
    if isinstance(o, dict):
        o = o.get('ad_entities', o)
    if isinstance(o, str):
        o = json.loads(o)
    return o or []


def money(x):
    if isinstance(x, dict):
        x = x.get('value')
    try:
        return float(x) if x not in (None, '') else 0.0
    except (TypeError, ValueError):
        return 0.0


def intv(x):
    try:
        return int(float(x)) if x not in (None, '') else 0
    except (TypeError, ValueError):
        return 0


# ---------- gerenciador (anúncio x dia) ----------
g = pd.DataFrame([dict(date=r['date_start'], campaign=r.get('campaign_name', ''), adset=r.get('adset_name', ''), ad=r.get('name', ''),
                       spend=money(r.get('amount_spent')), impressions=intv(r.get('impressions')), link_clicks=intv(r.get('link_click')),
                       page_view=intv(r.get('landing_page_view')), init_checkout=0, purchase=0, status=r.get('effective_status', ''))
                  for r in meta_rows('ads')])

# Campanhas duplicadas no Gerenciador mantêm a UTM antiga (ex.: "...-2026-07-24 — Cópia").
# Se o utm_term não bate com nenhuma campanha do Meta, casa pelo nome sem a data final.
_base = lambda n: re.sub(r'-\d{4}-\d{2}-\d{2}.*$', '', n or '').strip()
_por_base = collections.defaultdict(set)
for c in g['campaign'].unique():
    _por_base[_base(c)].add(c)
def campanha_meta(term):
    if not term or term in _por_base.get(_base(term), ()):
        return term
    cands = _por_base.get(_base(term), set())
    return next(iter(cands)) if len(cands) == 1 else term

# ---------- leads (planilha) ----------
import openpyxl
ws = openpyxl.load_workbook(f'{BASE}/leads.xlsx', read_only=True, data_only=True)['OCT26']
L, emails = [], []
for r in list(ws.iter_rows(values_only=True))[1:]:
    r = list(r or []) + [None] * 8
    if not B.iso(r[0]):
        continue
    nome, email = str(r[1] or '').strip(), str(r[2] or '').strip().lower()
    if TESTE.search(nome) or TESTE.search(email):
        continue
    src = str(r[3] or '').strip().lower()
    term = campanha_meta(str(r[5] or '').strip())
    tipo = 'HOT' if 'HOT' in term.upper() else 'COLD' if 'COLD' in term.upper() else 'OUTRO'
    row = dict(date=r[0], campaign=term, adset=str(r[4] or '').strip(), ad=str(r[6] or '').strip(),
               canal=src if src in PAGOS else 'sem_rastreio', tipo=tipo)
    L.append(row)
    emails.append(email)
ldf = pd.DataFrame(L)

# ---------- segmentos (Meta, por campanha e dia) ----------
def seg(rows, chave_fn):
    return pd.DataFrame([dict(date=r['date_start'], campaign=r.get('campaign_name', ''), chave=chave_fn(r),
                              spend=money(r.get('amount_spent')), resultado=intv(r.get('lead'))) for r in rows])

GEN = {'female': 'Feminino', 'male': 'Masculino'}
POS = {('instagram', 'feed'): 'IG Feed', ('instagram', 'instagram_stories'): 'IG Stories', ('instagram', 'instagram_reels'): 'IG Reels',
       ('instagram', 'instagram_explore'): 'IG Explorar', ('instagram', 'instagram_explore_grid_home'): 'IG Explorar',
       ('instagram', 'instagram_profile_feed'): 'IG Perfil', ('facebook', 'feed'): 'FB Feed', ('facebook', 'facebook_stories'): 'FB Stories',
       ('facebook', 'facebook_reels'): 'FB Reels', ('facebook', 'video_feeds'): 'FB Vídeo', ('facebook', 'marketplace'): 'FB Marketplace',
       ('audience_network', 'an_classic'): 'Audience Network', ('threads', 'threads_stream'): 'Threads'}
ig = meta_rows('idadegen')
segs = {
    'genero': seg(ig, lambda r: GEN.get(str(r.get('gender', '')).lower(), 'Desconhecido')),
    'idade': seg([r for r in ig if str(r.get('age', '')).lower() != 'unknown'], lambda r: r.get('age', '')),
    'posicao': seg(meta_rows('posicao'), lambda r: POS.get((r.get('publisher_platform', ''), r.get('platform_position', '')),
                                                          f"{r.get('publisher_platform', '')} {r.get('platform_position', '')}".strip())),
}

# ---------- lead scoring (mesmos critérios da rotina) ----------
def pts_comp(v):
    s = str(v or '').lower()
    return 3 if 'totalmente comprometid' in s else 2 if ('dudas' in s or 'quiero aprender' in s) else 1 if 'empezando a moverme' in s else 0
def pts_tempo(v):
    s = str(v or '').lower()
    return 3 if 'lo que sea necesario' in s else 2 if 'una hora a la semana' in s else 1 if 'un ratito' in s else 0
def pts_renda(v):
    s = str(v or '').strip().lower()
    if not s or 'menos de 1000' in s or 'menos de 1.000' in s: return 0
    if 'entre 1000 y 2000' in s or 'entre 1.000 y 2.000' in s: return 1
    return 2

rows2 = list(openpyxl.load_workbook(f'{BASE}/pesquisa.xlsx', read_only=True, data_only=True)['Respuestas de formulario 1'].iter_rows(values_only=True))
H = [str(h or '').lower() for h in rows2[0]]
col = lambda k, res: next((i for i, h in enumerate(H) if k in h), res)
cE, cC, cT, cR = col('correo', 1), col('compromiso', 15), col('dedicar a tu crecimiento', 18), col('ingresos', 23)
pesq = {}
for r in rows2[1:]:
    r = list(r or []) + [None] * 30
    if not r[0] or not r[cE]:
        continue
    e = str(r[cE]).strip().lower()
    if TESTE.search(e):
        continue
    p = pts_comp(r[cC]) + pts_tempo(r[cT]) + pts_renda(r[cR])
    pesq[e] = 'Ideal' if p >= 6 else 'Bom' if p >= 3 else 'Descarte'

vistos, grupos, tot = set(), collections.OrderedDict(), collections.Counter()
for l, e in zip(L, emails):
    if not e or e in vistos:
        continue
    vistos.add(e)
    sc = pesq.get(e)
    if not sc:
        tot['Sem dados'] += 1
        continue
    tot[sc] += 1
    k = (l['tipo'], l['campaign'] or 'sem_utm', l['adset'] or 'sem_utm', l['ad'] or 'sem_utm', l['canal'])
    gr = grupos.setdefault(k, dict(tipo=k[0], campanha=k[1], conjunto=k[2], criativo=k[3], canal=k[4], ideal=0, bom=0, descarte=0, total=0))
    gr[sc.lower()] += 1
    gr['total'] += 1

def expl(lg):
    pt = lg == 'pt'
    return dict(
        titulo='Como classificamos os leads' if pt else 'Cómo clasificamos los leads',
        intro=('Cruzamos cada cadastro com a pesquisa pelo e-mail e somamos pontos em 3 perguntas (máximo 8). '
               'Cada pessoa conta uma vez, com a UTM do primeiro cadastro.') if pt else
              ('Cruzamos cada registro con la encuesta por email y sumamos puntos en 3 preguntas (máximo 8). '
               'Cada persona cuenta una vez, con la UTM del primer registro.'),
        criterios=[
            dict(nome='Compromisso com mudanças' if pt else 'Compromiso con cambios', max='até 3 pts' if pt else 'hasta 3 pts',
                 regras=[['Totalmente comprometido/a', '3'], ['Dúvidas, mas quer aprender' if pt else 'Dudas, pero quiero aprender', '2'],
                         ['Começando a se mexer' if pt else 'Empezando a moverme', '1'], ['Custa se comprometer' if pt else 'Me cuesta comprometerme', '0']]),
            dict(nome='Tempo disponível' if pt else 'Tiempo disponible', max='até 3 pts' if pt else 'hasta 3 pts',
                 regras=[['O que for necessário' if pt else 'Lo que sea necesario', '3'], ['Uma hora por semana' if pt else 'Una hora a la semana', '2'],
                         ['Um pouco quando puder' if pt else 'Un ratito cuando pueda', '1'], ['Ainda não sabe' if pt else 'Aún no estoy seguro/a', '0']]),
            dict(nome='Renda familiar mensal' if pt else 'Ingresos familiares', max='até 2 pts' if pt else 'hasta 2 pts',
                 regras=[['2.000 € ou mais' if pt else '2.000 € o más', '2'], ['1.000 – 2.000 €', '1'], ['Menos de 1.000 €', '0']]),
        ],
        niveis=[['Ideal', 'Ideal', '6 a 8 pontos' if pt else '6 a 8 puntos'], ['Bom', 'Bom' if pt else 'Bueno', '3 a 5 pontos' if pt else '3 a 5 puntos'],
                ['Descarte', 'Descarte', '0 a 2 pontos' if pt else '0 a 2 puntos'],
                ['Sem', 'Sem pesquisa' if pt else 'Sin encuesta', 'Cadastrou mas ainda não respondeu' if pt else 'Se registró pero aún no respondió']])

lead_scoring = dict(
    totais={k: tot.get(k, 0) for k in ('Ideal', 'Bom', 'Descarte', 'Sem dados')},
    criterios={'Ideal': {'pt': '6 a 8 pts', 'es': '6 a 8 pts'}, 'Bom': {'pt': '3 a 5 pts', 'es': '3 a 5 pts'},
               'Descarte': {'pt': '0 a 2 pts', 'es': '0 a 2 pts'}},
    explicacao={'pt': expl('pt'), 'es': expl('es')},
    grupos=sorted(grupos.values(), key=lambda x: (-x['ideal'], -x['total'])))

try:
    leitura = json.load(open(f'{BASE}/leitura.json', encoding='utf-8'))
except FileNotFoundError:
    leitura = None

meta = dict(cliente='Familias Estelares', letra='FE', cor='#00a878', evento='PLF La cara oculta de la adolescencia',
            lancamento_cod='PLF-oct26', atualizado_em=AGORA.strftime('%d/%m/%Y %H:%M'),
            data_ate=(AGORA.date() - timedelta(days=1)).isoformat(), moeda='€', idioma_padrao='es', fonte_label='Planilha')
config = dict(modo='leads', paginas=['captacao', 'meta', 'metricas', 'scoring'], segmentos_estimado=False,
              metas={'ctr': [1.2, 0.9], 'connect_rate': [70, 60], 'tx_conv': [15, 10], 'cpm': [14, 20], 'cpl': [6, 7]})

d = B.montar(meta=meta, config=config, gerenciador=g, leads_df=ldf, segmentos=segs, lead_scoring=lead_scoring, leitura=leitura)
B.salvar(d, f'{BASE}/data-plf-oct26.json')
print('ok', f'{BASE}/data-plf-oct26.json', 'anuncios', len(d['anuncios']), 'leads', sum(r['n'] for r in d['leads']),
      'gasto', round(sum(r['sp'] for r in d['anuncios']), 2), 'scoring', lead_scoring['totais'])
